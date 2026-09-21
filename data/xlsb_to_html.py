"""
일일판매재고.xlsb → 단일 HTML 뷰어 생성.

- 시트별 값/서식(폰트·굵기·색·정렬·병합·컬럼폭·행높이·테두리·숨김) 재현
- 테마 색상 + tint 지원 (Office 기본 테마)
- 날짜 컬럼은 최근 7일치만 노출 (엑셀 날짜 시리얼 40000-60000 범위 자동 감지)
- 상단 탭 UI로 7개 시트 전환

의존: pywin32 (xlsb→xlsx 변환), openpyxl
"""
import os
import re
import sys
import io
import html
import datetime as dt
import tempfile
from pathlib import Path

# stdout UTF-8
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

BASE = Path(__file__).parent
SRC_XLSB = BASE / '일일판매재고.xlsb'
OUT_HTML = BASE / '일일판매재고_뷰어.html'

# 최근 N일 (오늘 포함)
DATE_WINDOW_DAYS = 7
EXCEL_EPOCH = dt.date(1899, 12, 30)
TODAY = dt.date.today()
MIN_KEEP_DATE = TODAY - dt.timedelta(days=DATE_WINDOW_DAYS - 1)
DATE_SERIAL_RANGE = (40000, 60000)


def xlsb_to_xlsx(src: Path) -> Path:
    """Excel COM 으로 xlsb → xlsx 변환. 모든 시트 unhide."""
    import win32com.client
    dst = Path(tempfile.gettempdir()) / (src.stem + '_conv.xlsx')
    if dst.exists():
        dst.unlink()
    excel = win32com.client.DispatchEx('Excel.Application')
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb = excel.Workbooks.Open(str(src), ReadOnly=True, UpdateLinks=0)
        for i in range(1, wb.Sheets.Count + 1):
            wb.Sheets(i).Visible = -1
        wb.SaveAs(str(dst), FileFormat=51)
        wb.Close(SaveChanges=False)
    finally:
        excel.Quit()
    return dst


def argb_to_css(argb) -> str | None:
    if not argb:
        return None
    s = str(argb)
    if not re.match(r'^[0-9A-Fa-f]{6,8}$', s):
        return None
    if len(s) == 8:
        return '#' + s[2:]
    if len(s) == 6:
        return '#' + s
    return None


# 셀 참조용 테마 인덱스 순서 (0=lt1, 1=dk1, 2=lt2, 3=dk2, 4..=accent1..6, 10=hlink, 11=folHlink)
THEME_CELL_ORDER = ['lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6', 'hlink', 'folHlink']
THEME_MAP: dict[int, str] = {}  # idx → 'RRGGBB'


def load_theme(wb):
    """Workbook.loaded_theme XML 에서 clrScheme 을 읽어 THEME_MAP 채움."""
    theme = wb.loaded_theme
    if isinstance(theme, bytes):
        theme = theme.decode('utf-8', errors='ignore')
    if not theme:
        return
    m = re.search(r'<a:clrScheme[^>]*>(.*?)</a:clrScheme>', theme, re.S)
    if not m:
        return
    scheme = m.group(1)
    xml_order = {}
    for match in re.finditer(r'<a:(\w+)>\s*<a:(srgbClr|sysClr)\s+([^/]+)/>', scheme):
        name = match.group(1)
        attrs = match.group(3)
        vm = re.search(r'val="([^"]+)"', attrs)
        lm = re.search(r'lastClr="([^"]+)"', attrs)
        color = (lm.group(1) if lm else (vm.group(1) if vm else '000000'))
        xml_order[name] = color.upper()
    for i, name in enumerate(THEME_CELL_ORDER):
        if name in xml_order:
            THEME_MAP[i] = xml_order[name]


def apply_tint(hex6: str, tint: float) -> str:
    """RGB 튜닝. tint > 0 : lighten, tint < 0 : darken."""
    if not tint:
        return hex6
    r = int(hex6[0:2], 16)
    g = int(hex6[2:4], 16)
    b = int(hex6[4:6], 16)
    if tint > 0:
        r = int(r + (255 - r) * tint)
        g = int(g + (255 - g) * tint)
        b = int(b + (255 - b) * tint)
    else:
        f = 1 + tint  # tint 는 음수
        r = int(r * f)
        g = int(g * f)
        b = int(b * f)
    r, g, b = max(0, min(255, r)), max(0, min(255, g)), max(0, min(255, b))
    return f'{r:02X}{g:02X}{b:02X}'


def color_to_css(color_obj) -> str | None:
    """openpyxl Color(rgb/theme/indexed) → '#rrggbb' 또는 None."""
    if color_obj is None:
        return None
    try:
        ctype = color_obj.type
    except Exception:
        ctype = None
    if ctype == 'rgb':
        v = color_obj.value
        if not v:
            return None
        s = str(v)
        if not re.match(r'^[0-9A-Fa-f]{6,8}$', s):
            return None
        hex6 = s[-6:] if len(s) >= 6 else s
        # 완전투명 (00000000) 은 None
        if len(s) == 8 and s[:2] == '00':
            return None
        tint = getattr(color_obj, 'tint', 0.0) or 0.0
        return '#' + apply_tint(hex6, tint)
    if ctype == 'theme':
        idx = color_obj.value
        base = THEME_MAP.get(idx)
        if not base:
            return None
        tint = getattr(color_obj, 'tint', 0.0) or 0.0
        return '#' + apply_tint(base, tint)
    return None


def fmt_date_md(d) -> str:
    """년 제거한 M/D (앞자리 0 제거)."""
    return f'{d.month}/{d.day}'


def is_date_format(nf: str) -> bool:
    nfl = (nf or '').lower()
    return any(tok in nfl for tok in ['yyyy', 'yy', 'm/d', 'd-m', 'mmm', 'mm-dd', 'mm/dd', 'dd/mm']) or 'm"/"d' in nfl


def fmt_value(v, number_format, decimals_override=None):
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, dt.datetime):
        return fmt_date_md(v.date())
    if isinstance(v, dt.date):
        return fmt_date_md(v)
    if isinstance(v, (int, float)):
        iv = int(v)
        if is_date_format(number_format) and DATE_SERIAL_RANGE[0] <= iv <= DATE_SERIAL_RANGE[1] and abs(v - iv) < 0.001:
            return fmt_date_md(EXCEL_EPOCH + dt.timedelta(days=iv))
        if decimals_override is not None:
            if decimals_override == 0:
                return f'{int(round(v)):,}'
            return f'{v:,.{decimals_override}f}'
        nf = (number_format or '').lower()
        if '%' in nf:
            s = f'{v*100:.2f}'.rstrip('0').rstrip('.')
            return s + '%'
        if '#,##0' in nf or ',' in nf:
            if isinstance(v, float) and not v.is_integer():
                return f'{v:,.2f}'
            return f'{int(v):,}'
        if isinstance(v, float):
            if v.is_integer():
                return str(int(v))
            return f'{v:.4f}'.rstrip('0').rstrip('.')
        return str(v)
    try:
        return v.strftime('%Y-%m-%d')
    except Exception:
        pass
    return str(v)


def border_side(side):
    if not side or not side.style:
        return None
    color = color_to_css(side.color) or '#000'
    style = side.style
    css_style = 'solid'
    if style in ('dotted',):
        css_style = 'dotted'
    elif style in ('dashed', 'dashDot', 'dashDotDot'):
        css_style = 'dashed'
    elif style in ('double',):
        css_style = 'double'
    width = '2px' if style in ('medium', 'thick', 'double') else '1px'
    return f'{width} {css_style} {color}'


def cell_style(cell):
    """openpyxl Cell → CSS style dict."""
    css = {}

    font = cell.font
    if font:
        if font.bold: css['font-weight'] = '700'
        if font.italic: css['font-style'] = 'italic'
        if font.underline: css['text-decoration'] = 'underline'
        if font.size: css['font-size'] = f'{float(font.size)}px'
        if font.name: css['font-family'] = f'"{font.name}", "맑은 고딕", sans-serif'
        col = color_to_css(font.color)
        if col: css['color'] = col

    fill = cell.fill
    if fill and fill.patternType == 'solid':
        bg = color_to_css(fill.fgColor)
        if bg: css['background-color'] = bg

    al = cell.alignment
    if al:
        if al.horizontal:
            css['text-align'] = al.horizontal
        if al.vertical:
            v = al.vertical
            css['vertical-align'] = {'top': 'top', 'bottom': 'bottom', 'center': 'middle'}.get(v, v)
        if al.wrap_text:
            css['white-space'] = 'pre-wrap'
            css['word-break'] = 'break-word'
        else:
            css['white-space'] = 'nowrap'

    b = cell.border
    if b:
        for side_name, css_prop in (('top','border-top'), ('bottom','border-bottom'), ('left','border-left'), ('right','border-right')):
            s = border_side(getattr(b, side_name, None))
            if s: css[css_prop] = s

    return css


def style_str(css_dict):
    return ';'.join(f'{k}:{v}' for k, v in css_dict.items())


def detect_date_columns(ws, max_col: int):
    """각 컬럼의 상단 10 행에서 날짜 셀(datetime 또는 엑셀 시리얼)을 찾음.
    반환: {col: date}"""
    result = {}
    scan_rows = min(10, ws.max_row or 1)
    for c in range(1, max_col + 1):
        for r in range(1, scan_rows + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, dt.datetime):
                result[c] = v.date()
                break
            if isinstance(v, dt.date):
                result[c] = v
                break
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                iv = int(v)
                if DATE_SERIAL_RANGE[0] <= iv <= DATE_SERIAL_RANGE[1] and abs(v - iv) < 0.001:
                    result[c] = EXCEL_EPOCH + dt.timedelta(days=iv)
                    break
    return result


def column_has_data(ws, c: int, header_end_row: int = 10) -> bool:
    """헤더 아래 데이터 행(header_end_row+1..max_row)에 어떤 값이라도 있는지."""
    for r in range(header_end_row + 1, (ws.max_row or 0) + 1):
        v = ws.cell(row=r, column=c).value
        if v not in (None, '', 0):
            return True
    return False


def filter_date_columns(ws, date_cols: dict, today) -> set[int]:
    """미래·데이터없는 컬럼 숨기고, 데이터 있는 가장 최근 7일만 유지.
    반환: 숨길 컬럼 index set."""
    if not date_cols:
        return set()
    keep_candidates = []  # [(date, col)]
    hide_because_future_or_empty = set()
    for c, d in date_cols.items():
        if d > today:
            hide_because_future_or_empty.add(c)
            continue
        if not column_has_data(ws, c):
            hide_because_future_or_empty.add(c)
            continue
        keep_candidates.append((d, c))
    keep_candidates.sort(reverse=True)  # 최신 → 오래된
    keep_cols = {c for _, c in keep_candidates[:DATE_WINDOW_DAYS]}
    hidden = set(date_cols.keys()) - keep_cols
    return hidden


def detect_header_row(ws, max_col: int, max_row: int) -> int:
    """상단 15행 중 텍스트 셀이 가장 많이 채워진 행을 헤더 행으로 판단."""
    best_r, best_count = 1, -1
    for r in range(1, min(16, max_row + 1)):
        cnt = 0
        for c in range(1, max_col + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip():
                cnt += 1
        if cnt > best_count:
            best_count = cnt
            best_r = r
    return best_r


def render_sheet(ws, style_cache: dict, prefix: str, sheet_idx: int):
    """openpyxl Worksheet → <table> HTML. 반복 스타일은 클래스로 압축."""
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1
    header_row = detect_header_row(ws, max_col, max_row)
    data_start_row = header_row + 1

    def get_class(css_str):
        if not css_str:
            return ''
        if css_str in style_cache:
            return style_cache[css_str]
        cls = f'{prefix}{len(style_cache)}'
        style_cache[css_str] = cls
        return cls

    # 날짜 컬럼 감지 → (미래 제외 + 데이터 있는) 최근 7일만
    date_cols = detect_date_columns(ws, max_col)
    hidden_date_cols = filter_date_columns(ws, date_cols, TODAY)

    # 컬럼 헤더 텍스트 수집
    col_headers: dict[int, str] = {}
    for c in range(1, max_col + 1):
        parts = []
        for r in range(1, min(header_row + 1, max_row + 1)):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip():
                parts.append(v.strip())
        col_headers[c] = ' '.join(parts)

    # 필터 가능한 텍스트 컬럼 감지 (헤더 아래 데이터가 텍스트, 유니크 2~30개)
    filterable_cols: dict[int, list[str]] = {}
    for c in range(1, max_col + 1):
        if c in hidden_date_cols:
            continue
        vals = set()
        text_count = 0
        total_count = 0
        for r in range(data_start_row, min(data_start_row + 1000, max_row + 1)):
            v = ws.cell(row=r, column=c).value
            if v is None or v == '':
                continue
            total_count += 1
            if isinstance(v, str):
                s = v.strip()
                if s:
                    vals.add(s)
                    text_count += 1
                    if len(vals) > 50:
                        break
        if 2 <= len(vals) <= 30 and text_count >= max(3, total_count * 0.5):
            filterable_cols[c] = sorted(vals)

    def decimals_for(c: int) -> int | None:
        h = col_headers.get(c, '')
        if '판매가능일수' in h:
            return 1
        if '평균판매량' in h:
            return 0
        if '공급가' in h or '원가' in h or '단가' in h or '판매가' in h or '금액' in h:
            return 0
        return None

    merged = {}
    for r in ws.merged_cells.ranges:
        rs = r.max_row - r.min_row + 1
        cs = r.max_col - r.min_col + 1
        merged[(r.min_row, r.min_col)] = (rs, cs)
        for rr in range(r.min_row, r.max_row + 1):
            for cc in range(r.min_col, r.max_col + 1):
                if (rr, cc) != (r.min_row, r.min_col):
                    merged[(rr, cc)] = 'skip'

    hidden_cols: set[int] = set(hidden_date_cols)
    for c in range(1, max_col + 1):
        letter = ws.cell(row=1, column=c).column_letter
        cd = ws.column_dimensions.get(letter)
        if cd and cd.hidden:
            hidden_cols.add(c)

    colgroup = ['<colgroup>']
    for c in range(1, max_col + 1):
        if c in hidden_cols:
            continue
        letter = ws.cell(row=1, column=c).column_letter
        cd = ws.column_dimensions.get(letter)
        w = float(cd.width) * 7.5 if (cd and cd.width) else 64
        colgroup.append(f'<col style="width:{w:.0f}px">')
    colgroup.append('</colgroup>')

    rows_html = []
    for r in range(1, max_row + 1):
        rd = ws.row_dimensions.get(r)
        if rd and rd.hidden:
            continue
        h = (rd.height * 1.33) if (rd and rd.height) else None
        row_style = f'height:{h:.0f}px' if h else ''
        is_data = r >= data_start_row
        row_attrs = f' data-r="{r}"'
        row_cls = ' class="data-row"' if is_data else ''
        cells = []
        for c in range(1, max_col + 1):
            if c in hidden_cols:
                continue
            key = (r, c)
            if merged.get(key) == 'skip':
                continue
            spans = merged.get(key)
            span_attr = ''
            if isinstance(spans, tuple):
                rs, cs = spans
                if cs > 1:
                    visible_cs = sum(1 for cc in range(c, c + cs) if cc not in hidden_cols)
                    cs = visible_cs
                if rs > 1: span_attr += f' rowspan="{rs}"'
                if cs > 1: span_attr += f' colspan="{cs}"'
            cell = ws.cell(row=r, column=c)
            nf = cell.number_format
            if c in date_cols and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool) \
                    and DATE_SERIAL_RANGE[0] <= int(cell.value) <= DATE_SERIAL_RANGE[1]:
                nf = 'm/d'
            v = fmt_value(cell.value, nf, decimals_override=decimals_for(c))
            css = style_str(cell_style(cell))
            cls_name = get_class(css)
            base_cls = cls_name
            cls_extra = ''
            attr_extra = f' data-c="{c}"'
            # 필터 대상 컬럼의 원본 값 저장 (텍스트)
            if is_data and c in filterable_cols and isinstance(cell.value, str):
                attr_extra += f' data-v="{html.escape(cell.value.strip(), quote=True)}"'
            # 헤더 셀에 필터 버튼
            filter_btn = ''
            if r == header_row and c in filterable_cols:
                filter_btn = f'<span class="flt-btn" data-col="{c}" title="필터">▾</span>'
            # 편집 가능
            edit_attr = ' contenteditable="true"' if is_data else ''
            cls_attr = f' class="{base_cls}{cls_extra}"' if base_cls else (f' class="{cls_extra.strip()}"' if cls_extra else '')
            cells.append(f'<td{span_attr}{cls_attr}{attr_extra}{edit_attr}>{html.escape(v)}{filter_btn}</td>')
        style_attr = f' style="{row_style}"' if row_style else ''
        rows_html.append(f'<tr{row_cls}{row_attrs}{style_attr}>{"".join(cells)}</tr>')

    tbl = f'<table class="sheet-table">{"".join(colgroup)}<tbody>{"".join(rows_html)}</tbody></table>'
    filter_meta = {str(c): {'header': col_headers.get(c, ''), 'values': vals} for c, vals in filterable_cols.items()}
    return tbl, filter_meta


def build_html(sheets_html, style_cache: dict) -> str:
    """탭 UI 로 감싸서 단일 HTML 반환. sheets_html: list of (name, tbl_html, filter_meta)"""
    import json
    tabs = []
    panels = []
    filter_configs = {}
    for i, (name, tbl, fmeta) in enumerate(sheets_html):
        cls = 'tab active' if i == 0 else 'tab'
        tabs.append(f'<button class="{cls}" data-idx="{i}">{html.escape(name)}</button>')
        pcls = 'panel active' if i == 0 else 'panel'
        panels.append(f'<div class="{pcls}" data-idx="{i}">{tbl}</div>')
        filter_configs[i] = fmeta
    style_rules = '\n'.join(f'.{cls}{{{css}}}' for css, cls in style_cache.items())
    filter_json = json.dumps(filter_configs, ensure_ascii=False)
    return f'''<!doctype html>
<html lang="ko">
<head>
<meta charset="utf-8">
<title>일일판매재고 뷰어</title>
<style>
  * {{ box-sizing: border-box; }}
  html, body {{ margin: 0; height: 100%; font-family: "맑은 고딕", "Malgun Gothic", sans-serif; font-size: 12px; color: #222; background: #f4f4f4; }}
  .tabs {{ position: sticky; top: 0; z-index: 10; background: #f0efe9; border-bottom: 1px solid #c8c8c8; padding: 4px 6px 0; display: flex; gap: 2px; overflow-x: auto; }}
  .tab {{ background: #e6e4dd; border: 1px solid #c8c8c8; border-bottom: none; border-radius: 4px 4px 0 0; padding: 6px 14px; cursor: pointer; font-size: 12px; white-space: nowrap; color: #333; }}
  .tab.active {{ background: #fff; font-weight: 700; padding-bottom: 7px; margin-bottom: -1px; }}
  .tab:hover:not(.active) {{ background: #efeee7; }}
  .panels {{ padding: 8px; }}
  .panel {{ display: none; background: #fff; border: 1px solid #d9d9d9; overflow: auto; max-height: calc(100vh - 60px); }}
  .panel.active {{ display: block; }}
  table.sheet-table {{ border-collapse: collapse; font-size: 11px; table-layout: fixed; }}
  table.sheet-table td {{ padding: 1px 4px; border: 1px solid #d9d9d9; overflow: hidden; text-overflow: ellipsis; }}
  td[contenteditable="true"]:focus {{ outline: 2px solid #2563eb; background: #fffbea; }}
  .flt-btn {{ display: inline-block; margin-left: 4px; padding: 0 3px; font-size: 10px; color: #666; cursor: pointer; user-select: none; border: 1px solid transparent; border-radius: 2px; }}
  .flt-btn:hover {{ background: #ffe58a; border-color: #d0b040; }}
  .flt-btn.active {{ background: #ffc000; color: #000; border-color: #806000; }}
  .flt-popup {{ position: fixed; z-index: 100; background: #fff; border: 1px solid #999; box-shadow: 0 4px 16px rgba(0,0,0,.15); padding: 8px; max-height: 300px; overflow-y: auto; min-width: 180px; font-size: 12px; }}
  .flt-popup .flt-head {{ display: flex; gap: 6px; margin-bottom: 6px; }}
  .flt-popup input[type="search"] {{ flex: 1; padding: 3px 6px; border: 1px solid #bbb; border-radius: 2px; }}
  .flt-popup .flt-actions {{ display: flex; gap: 4px; margin-bottom: 6px; }}
  .flt-popup .flt-actions button {{ font-size: 11px; padding: 2px 8px; cursor: pointer; }}
  .flt-popup label {{ display: flex; gap: 6px; padding: 2px 0; cursor: pointer; }}
  .flt-popup label:hover {{ background: #f2f2f2; }}
  .flt-popup .flt-apply {{ margin-top: 6px; padding: 4px 8px; background: #2563eb; color: #fff; border: none; border-radius: 3px; cursor: pointer; width: 100%; }}
  {style_rules}
</style>
</head>
<body>
<div class="tabs">{''.join(tabs)}</div>
<div class="panels">{''.join(panels)}</div>
<div class="flt-popup" id="fltPopup" style="display:none"></div>
<script>
const FILTER_CFG = {filter_json};
// 각 시트별 활성 필터: {{sheetIdx: {{col: Set(selectedValues)}}}}
const ACTIVE = {{}};

document.querySelectorAll('.tab').forEach(t => {{
  t.addEventListener('click', () => {{
    const i = t.dataset.idx;
    document.querySelectorAll('.tab').forEach(x => x.classList.toggle('active', x.dataset.idx === i));
    document.querySelectorAll('.panel').forEach(x => x.classList.toggle('active', x.dataset.idx === i));
    document.getElementById('fltPopup').style.display = 'none';
  }});
}});

document.addEventListener('click', (ev) => {{
  const btn = ev.target.closest('.flt-btn');
  const popup = document.getElementById('fltPopup');
  if (btn) {{
    ev.stopPropagation();
    const col = btn.dataset.col;
    const panel = btn.closest('.panel');
    const sheetIdx = panel.dataset.idx;
    const cfg = (FILTER_CFG[sheetIdx] || {{}})[col];
    if (!cfg) return;
    openFilterPopup(btn, sheetIdx, col, cfg);
    return;
  }}
  if (!ev.target.closest('.flt-popup')) popup.style.display = 'none';
}});

function openFilterPopup(anchor, sheetIdx, col, cfg) {{
  const popup = document.getElementById('fltPopup');
  const active = (ACTIVE[sheetIdx] && ACTIVE[sheetIdx][col]) || null;
  const escapeAttr = s => s.replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;').replace(/"/g,'&quot;');
  const checks = cfg.values.map((v, i) => {{
    const checked = !active || active.has(v);
    return `<label><input type="checkbox" data-v="${{escapeAttr(v)}}" ${{checked?'checked':''}}> <span>${{escapeAttr(v)}}</span></label>`;
  }}).join('');
  popup.innerHTML = `
    <div class="flt-head"><b>${{escapeAttr(cfg.header || '')}}</b></div>
    <div class="flt-actions">
      <button data-act="all">전체</button>
      <button data-act="none">해제</button>
    </div>
    <input type="search" placeholder="검색" style="width:100%;margin-bottom:6px">
    <div class="flt-list">${{checks}}</div>
    <button class="flt-apply">적용</button>
  `;
  const r = anchor.getBoundingClientRect();
  popup.style.left = Math.min(window.innerWidth - 220, r.left) + 'px';
  popup.style.top = (r.bottom + 4) + 'px';
  popup.style.display = 'block';
  popup.querySelector('input[type="search"]').addEventListener('input', (e) => {{
    const q = e.target.value.toLowerCase();
    popup.querySelectorAll('.flt-list label').forEach(l => {{
      l.style.display = l.textContent.toLowerCase().includes(q) ? '' : 'none';
    }});
  }});
  popup.querySelector('[data-act="all"]').addEventListener('click', () => {{
    popup.querySelectorAll('input[type="checkbox"]').forEach(cb => cb.checked = true);
  }});
  popup.querySelector('[data-act="none"]').addEventListener('click', () => {{
    popup.querySelectorAll('input[type="checkbox"]').forEach(cb => cb.checked = false);
  }});
  popup.querySelector('.flt-apply').addEventListener('click', () => {{
    const selected = new Set();
    popup.querySelectorAll('input[type="checkbox"]:checked').forEach(cb => selected.add(cb.dataset.v));
    ACTIVE[sheetIdx] = ACTIVE[sheetIdx] || {{}};
    // 전체 선택이면 필터 해제
    if (selected.size === cfg.values.length) delete ACTIVE[sheetIdx][col];
    else ACTIVE[sheetIdx][col] = selected;
    applyFilters(sheetIdx);
    popup.style.display = 'none';
  }});
}}

function applyFilters(sheetIdx) {{
  const panel = document.querySelector('.panel[data-idx="' + sheetIdx + '"]');
  const filters = ACTIVE[sheetIdx] || {{}};
  const cols = Object.keys(filters);
  // 필터 버튼에 activate 표시
  panel.querySelectorAll('.flt-btn').forEach(b => {{
    b.classList.toggle('active', filters.hasOwnProperty(b.dataset.col));
  }});
  panel.querySelectorAll('tr.data-row').forEach(tr => {{
    let show = true;
    for (const c of cols) {{
      const sel = filters[c];
      const td = tr.querySelector('td[data-c="' + c + '"]');
      const v = td ? (td.getAttribute('data-v') || td.textContent.trim()) : '';
      if (!sel.has(v)) {{ show = false; break; }}
    }}
    tr.style.display = show ? '' : 'none';
  }});
}}
</script>
</body>
</html>'''


def main():
    import openpyxl
    print(f'변환 중: {SRC_XLSB}')
    xlsx = xlsb_to_xlsx(SRC_XLSB)
    print(f'  → {xlsx} ({os.path.getsize(xlsx):,} bytes)')

    print(f'로드 중 (openpyxl)...')
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    load_theme(wb)
    print(f'  테마 색상 {len(THEME_MAP)}개 로드')
    sheets_html = []
    style_cache = {}
    for i, name in enumerate(wb.sheetnames):
        ws = wb[name]
        print(f'  렌더: {name}  ({ws.max_row} x {ws.max_column})')
        tbl, fmeta = render_sheet(ws, style_cache, prefix=f's{i}_', sheet_idx=i)
        sheets_html.append((name, tbl, fmeta))
    print(f'  고유 스타일 수: {len(style_cache)}')

    html_out = build_html(sheets_html, style_cache)
    OUT_HTML.write_text(html_out, encoding='utf-8')
    print(f'완료: {OUT_HTML}  ({os.path.getsize(OUT_HTML):,} bytes)')


if __name__ == '__main__':
    main()
