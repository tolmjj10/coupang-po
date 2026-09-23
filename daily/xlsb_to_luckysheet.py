"""
일일판매재고.xlsb + data/daily/*.csv → Luckysheet 뷰어(HTML) 생성.

- xlsb → xlsx (Excel COM)
- data/daily/*.csv 자동 집계 → 9월 출고 시트 해당 날짜 컬럼에 반영 (SKU별 출고수량 합)
- xlsx 전처리:
  * 최근 7일 데이터 있는 날짜 컬럼만 남기고 나머지 숨김
  * 특정 컬럼 서식 강제 적용
- 브라우저에서 LuckyExcel + Luckysheet 로 편집/수식 지원
"""
import os, sys, io, base64, tempfile, shutil, re, csv, datetime as dt
from pathlib import Path
from collections import defaultdict

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', write_through=True)

BASE = Path(__file__).parent
SRC_XLSB = BASE / '일일판매재고.xlsb'
CSV_DIR = BASE.parent / 'data' / 'daily'   # data/daily/*.csv 를 매일 합침
# 이 폴더 자체가 웹 배포 경로 (coupang-po.pages.dev/daily/)
OUT_HTML = BASE / 'index.html'

DATE_WINDOW_DAYS = 7
EXCEL_EPOCH = dt.date(1899, 12, 30)
TODAY = dt.date.today()
DATE_SERIAL_RANGE = (40000, 60000)

# 시트별 항상 숨길 컬럼
# B: 광고, C: 품목, D: 쿠팡카테고리, F: 광고취합용, P: 현재고*원가, Q: 현재고*공급가(+VAT)
ALWAYS_HIDDEN = {
    '9월 출고': {2, 3, 4, 6, 16, 17},
    '9월 입고': {2, 3, 4, 6, 16, 17},
}

# 마진테이블 시트: SKU# 컬럼 B, 조회 대상 범위 B:O
MARGIN_SHEET = '마진테이블'
MARGIN_LOOKUP_RANGE = "'마진테이블'!$B:$O"
# B:O 범위 내에서 각 컬럼 위치 (B=1)
MARGIN_COL_COST   = 10  # K: 원가
MARGIN_COL_SUPPLY = 5   # F: 공급가
MARGIN_COL_VAT    = 6   # G: 부가세
MARGIN_COL_MARGIN = 13  # N: 마진(-vat)
MARGIN_COL_RATE   = 14  # O: 이익율


IS_WINDOWS = sys.platform.startswith('win')


def xlsb_to_xlsx(src: Path) -> Path:
    """Windows: Excel COM. Linux: LibreOffice headless."""
    dst = Path(tempfile.gettempdir()) / (src.stem + '_conv.xlsx')
    if dst.exists():
        dst.unlink()
    if IS_WINDOWS:
        import win32com.client
        excel = win32com.client.DispatchEx('Excel.Application')
        excel.Visible = False; excel.DisplayAlerts = False
        try:
            wb = excel.Workbooks.Open(str(src), ReadOnly=True, UpdateLinks=0)
            for i in range(1, wb.Sheets.Count + 1):
                wb.Sheets(i).Visible = -1
            wb.SaveAs(str(dst), FileFormat=51)
            wb.Close(SaveChanges=False)
        finally:
            excel.Quit()
    else:
        import subprocess
        out_dir = Path(tempfile.gettempdir())
        subprocess.run(
            ['libreoffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', str(out_dir), str(src)],
            check=True, timeout=300,
        )
        conv = out_dir / (src.stem + '.xlsx')
        if conv != dst:
            shutil.move(str(conv), str(dst))
    return dst


def excel_recalc(path: Path) -> None:
    """수식 결과값을 캐시로 저장. Windows: Excel. Linux: LibreOffice."""
    if IS_WINDOWS:
        import win32com.client
        excel = win32com.client.DispatchEx('Excel.Application')
        excel.Visible = False; excel.DisplayAlerts = False
        try:
            wb = excel.Workbooks.Open(str(path), UpdateLinks=0)
            excel.CalculateFull()
            wb.Save()
            wb.Close(SaveChanges=False)
        finally:
            excel.Quit()
    else:
        import subprocess
        out_dir = path.parent
        # 임시 폴더에 변환 → 원본 덮어쓰기
        tmp_out = Path(tempfile.gettempdir()) / '_recalc_out'
        tmp_out.mkdir(parents=True, exist_ok=True)
        subprocess.run(
            ['libreoffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', str(tmp_out), str(path)],
            check=True, timeout=300,
        )
        result = tmp_out / (path.stem + '.xlsx')
        shutil.move(str(result), str(path))


def detect_date_columns(ws, max_col: int) -> dict:
    result = {}
    scan_rows = min(10, ws.max_row or 1)
    for c in range(1, max_col + 1):
        for r in range(1, scan_rows + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, dt.datetime):
                result[c] = v.date()
                break
            if isinstance(v, dt.date):
                result[c] = v
                break
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                iv = int(v)
                if DATE_SERIAL_RANGE[0] <= iv <= DATE_SERIAL_RANGE[1] and abs(v - iv) < 0.001:
                    result[c] = EXCEL_EPOCH + dt.timedelta(days=iv)
                    break
    return result


def column_has_data(ws, c: int, header_end_row: int = 10) -> bool:
    for r in range(header_end_row + 1, (ws.max_row or 0) + 1):
        v = ws.cell(row=r, column=c).value
        if v not in (None, '', 0):
            return True
    return False


def compute_hidden_date_cols(ws, date_cols: dict) -> set:
    if not date_cols:
        return set()
    keep = []
    for c, d in date_cols.items():
        if d > TODAY:
            continue
        if not column_has_data(ws, c):
            continue
        keep.append((d, c))
    keep.sort(reverse=True)
    keep_set = {c for _, c in keep[:DATE_WINDOW_DAYS]}
    return set(date_cols.keys()) - keep_set


def collect_headers(ws, max_col: int, up_to_row: int = 10) -> dict:
    headers = {}
    for c in range(1, max_col + 1):
        parts = []
        for r in range(1, min(up_to_row + 1, (ws.max_row or 0) + 1)):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip():
                parts.append(v.strip())
        headers[c] = ' '.join(parts)
    return headers


def decimals_format_for(header: str) -> str | None:
    if '판매가능일수' in header:
        return '#,##0.0'
    if '평균판매량' in header:
        return '#,##0'
    if any(k in header for k in ('공급가', '원가', '단가', '판매가', '금액')):
        return '#,##0'
    return None


def col_letter(c: int) -> str:
    """1-based col index → Excel letter."""
    s = ''
    while c > 0:
        c, rem = divmod(c - 1, 26)
        s = chr(65 + rem) + s
    return s


def parse_daily_csv(csv_path: Path) -> tuple[dt.date | None, dict]:
    """CSV 파싱 → (date, {sku: {'out','in','stock_fcvf'}}).
    stock_fcvf: 센터가 FC* 또는 VF* 로 시작하는 행들의 현재재고수량 합계."""
    per_sku = defaultdict(lambda: {'out': 0, 'in': 0, 'stock_fcvf': 0})
    date = None
    try:
        with open(csv_path, encoding='utf-8-sig', newline='') as f:
            reader = csv.DictReader(f)
            for row in reader:
                if not date:
                    d = (row.get('날짜') or '').strip()
                    try: date = dt.datetime.strptime(d, '%Y%m%d').date()
                    except ValueError: pass
                sku = (row.get('SKU ID') or '').strip()
                if not sku: continue
                center = (row.get('센터') or '').strip().upper()
                try:
                    per_sku[sku]['out'] += int(row.get('출고수량') or 0)
                    per_sku[sku]['in']  += int(row.get('입고수량') or 0)
                    if center.startswith('FC') or center.startswith('VF'):
                        per_sku[sku]['stock_fcvf'] += int(row.get('현재재고수량') or 0)
                except ValueError:
                    pass
    except Exception as e:
        print(f'  !! CSV 파싱 실패 {csv_path.name}: {e}')
    return date, dict(per_sku)


def aggregate_csvs_into_wb(wb, wb_vals):
    """data/daily/*.csv 를 읽어 9월 출고 / 9월 입고 시트 해당 날짜 컬럼에 반영."""
    if not CSV_DIR.exists():
        print(f'  CSV 폴더 없음 (스킵): {CSV_DIR}')
        return
    csvs = sorted(CSV_DIR.glob('basic_operation_rocket_*.csv'))
    if not csvs:
        print(f'  CSV 파일 없음: {CSV_DIR}')
        return

    per_date = {}
    for cp in csvs:
        d, data = parse_daily_csv(cp)
        if d: per_date[d] = data
    if not per_date:
        print('  CSV 유효 데이터 없음')
        return
    print(f'  CSV 집계: {len(per_date)}일 ({min(per_date)} ~ {max(per_date)})')

    latest_date = max(per_date)
    latest_stock = {sku: d['stock_fcvf'] for sku, d in per_date[latest_date].items()}

    for sheet_name, out_field in (('9월 출고', 'out'), ('9월 입고', 'in')):
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        ws_v = wb_vals[sheet_name]
        max_col = ws.max_column or 1
        max_row = ws.max_row or 1

        # SKU ID 컬럼 + 헤더 행 탐지
        sku_col = None
        header_row = None
        for r in range(1, min(15, max_row + 1)):
            for c in range(1, max_col + 1):
                v = ws_v.cell(row=r, column=c).value
                if isinstance(v, str) and v.strip() == 'SKU ID':
                    sku_col = c; header_row = r
                    break
            if sku_col: break
        if not sku_col:
            print(f'  [{sheet_name}] SKU ID 컬럼 못 찾음')
            continue

        # 현재고 컬럼 (헤더가 정확히 '현재고')
        stock_col = None
        for c in range(1, max_col + 1):
            v = ws_v.cell(row=header_row, column=c).value
            if isinstance(v, str) and v.strip() == '현재고':
                stock_col = c
                break

        # 날짜 → 컬럼 매핑
        date_to_col = {}
        for c in range(1, max_col + 1):
            v = ws_v.cell(row=header_row, column=c).value
            if isinstance(v, dt.datetime): date_to_col[v.date()] = c
            elif isinstance(v, dt.date):   date_to_col[v] = c

        # SKU → 행 매핑
        sku_to_row = {}
        for r in range(header_row + 1, max_row + 1):
            v = ws_v.cell(row=r, column=sku_col).value
            if v is None: continue
            sku = str(v).strip()
            if sku and sku not in sku_to_row:
                sku_to_row[sku] = r

        written = 0; missing_dates = 0; missing_skus = set()
        for date, sku_data in per_date.items():
            col = date_to_col.get(date)
            if not col: missing_dates += 1; continue
            for sku, data in sku_data.items():
                row = sku_to_row.get(sku)
                if not row: missing_skus.add(sku); continue
                ws.cell(row=row, column=col).value = data[out_field]
                written += 1
        print(f'  [{sheet_name}] 반영 {written:,}셀 · 날짜미매칭 {missing_dates} · SKU미매칭 {len(missing_skus)}')

        # 최신 CSV FC+VF 합계를 현재고 컬럼에 값으로 기록 (9월 출고 시트만)
        if stock_col and sheet_name == '9월 출고':
            stock_written = 0
            for sku, qty in latest_stock.items():
                row = sku_to_row.get(sku)
                if not row: continue
                ws.cell(row=row, column=stock_col).value = qty
                stock_written += 1
            print(f'  [{sheet_name}] 현재고(FC+VF) {stock_written:,}행 · 기준일 {latest_date}')


def preprocess_xlsx(src: Path) -> Path:
    """xlsx 전처리:
    - 수식 → 캐시된 값으로 대체 (Luckysheet formula.js 파싱 오류 회피)
    - 옛 날짜 컬럼 hidden=True
    - 지정 컬럼 서식 override
    - 주요 요약 컬럼(7일/14일/8월 평균판매량, 판매가능일수)에는 =AVERAGE / 나눗셈 수식 재삽입
      → 사용자가 날짜 셀 값을 바꾸면 자동 재계산됨
    """
    import openpyxl
    dst = Path(tempfile.gettempdir()) / '일일판매재고_prep.xlsx'
    shutil.copy(src, dst)

    wb_vals = openpyxl.load_workbook(dst, data_only=True)
    wb = openpyxl.load_workbook(dst)

    # data/daily/*.csv 를 시트에 반영 (수식 변환 전에 실행 → 이후 Excel 재계산이 값 반영)
    print('CSV 자동 반영 중...')
    aggregate_csvs_into_wb(wb, wb_vals)

    # CSV 반영 후 wb 를 값 기준으로 다시 읽어 hidden 판정용으로 사용
    # (하지만 파일에 아직 저장 전이므로 wb 자체를 직접 참조)
    # openpyxl 은 .value 가 cell.value 로 노출되므로 wb 그대로 has_data 판정 가능

    for name in wb.sheetnames:
        ws = wb[name]
        ws_v = wb_vals[name]
        max_col = ws.max_column or 1
        max_row = ws.max_row or 1
        date_cols = detect_date_columns(ws_v, max_col)
        # CSV 반영 후의 데이터로 판정하려면 ws(값이 최신) 사용
        hidden = compute_hidden_date_cols(ws, date_cols)
        # 시트별 항상 숨길 컬럼 추가
        hidden |= ALWAYS_HIDDEN.get(name, set())
        headers = collect_headers(ws_v, max_col)

        # 수식 → 값
        formula_replaced = 0
        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                cell = ws.cell(row=r, column=c)
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    v = ws_v.cell(row=r, column=c).value
                    cell.value = v
                    formula_replaced += 1

        # 헤더 행 찾기 (텍스트가 가장 많은 상단 15행)
        header_row = 1
        best = -1
        for r in range(1, min(16, max_row + 1)):
            cnt = sum(1 for c in range(1, max_col + 1) if isinstance(ws_v.cell(row=r, column=c).value, str))
            if cnt > best:
                best = cnt
                header_row = r
        data_start = header_row + 1

        # 데이터 있는 날짜 컬럼 (오늘 이하) 오름차순
        dated_with_data = sorted(
            [(d, c) for c, d in date_cols.items()
             if d <= TODAY and column_has_data(ws_v, c)]
        )
        # 마지막 N 개
        last7_cols = [c for _, c in dated_with_data[-7:]]
        last14_cols = [c for _, c in dated_with_data[-14:]]
        # 8월 판매량용: 8월 날짜만
        aug_cols = [c for d, c in dated_with_data if d.month == 8]
        # 9월 총합계용: 9월 날짜만
        sep_cols = [c for d, c in dated_with_data if d.month == 9]

        # 헤더 정규화 (공백/개행 제거)
        def norm(s: str) -> str:
            return re.sub(r'\s+', '', s or '')

        # 요약 컬럼 찾기 (정규화 매칭)
        col_targets = {}  # col idx → formula_type
        cost_col = supply_vat_col = supply_novat_col = margin_col = None
        seven_avg_col = stock_col = totalrev_col = None
        for c in range(1, max_col + 1):
            h = headers.get(c, '')
            n = norm(h)
            if '7일평균판매량' in n:
                col_targets[c] = '7avg'; seven_avg_col = c
            elif '14일평균판매량' in n:
                col_targets[c] = '14avg'
            elif '8월평균판매량' in n:
                col_targets[c] = '8avg'
            elif '8월판매량' in n:
                col_targets[c] = '8sum'
            elif '판매가능일수' in n:
                col_targets[c] = 'days'
            elif n == '현재고':
                stock_col = c
            elif '총합계' in n:
                col_targets[c] = 'sepSum'
            elif n == '원가':
                col_targets[c] = 'cost'; cost_col = c
            elif n == '원가총액':
                col_targets[c] = 'costTotal'
            elif n == '공급가(+VAT)':
                col_targets[c] = 'supplyVat'; supply_vat_col = c
            elif n == '공급가(-VAT)':
                col_targets[c] = 'supplyNoVat'; supply_novat_col = c
            elif n == '오아마진(-vat)':
                col_targets[c] = 'margin'; margin_col = c
            elif n == '오아총마진':
                col_targets[c] = 'marginTotal'
            elif n == '오아마진율':
                col_targets[c] = 'marginRate'
            elif n == '일일매출액':
                col_targets[c] = 'dailyRev'
            elif n == '총매출':
                col_targets[c] = 'totalRev'; totalrev_col = c
            elif n == '평균일매출':
                col_targets[c] = 'avgDailyRev'

        formulas_added = 0

        def range_or_list(cols, row):
            """연속된 컬럼이면 A:B 범위, 아니면 콤마 리스트."""
            if not cols:
                return None
            s = sorted(cols)
            if s[-1] - s[0] == len(s) - 1:
                return f'{col_letter(s[0])}{row}:{col_letter(s[-1])}{row}'
            return ','.join(f'{col_letter(cc)}{row}' for cc in s)

        # SKU ID 컬럼 (VLOOKUP 조회 키). 헤더에 'SKUID' 포함 매칭.
        sku_id_col = None
        for c in range(1, max_col + 1):
            n = norm(headers.get(c, ''))
            if 'SKUID' in n:
                sku_id_col = c
                break

        # 마진테이블 시트 존재 확인 (없으면 VLOOKUP 스킵)
        has_margin = MARGIN_SHEET in wb.sheetnames

        # 총합계 컬럼 (BJ) - marginTotal/costTotal 등에서 참조
        sep_sum_col = next((c for c, k in col_targets.items() if k == 'sepSum'), None)

        # 9월 date columns 중 데이터가 있는 마지막 컬럼(=전일자) & 데이터 일수
        sep_data_cols = sorted(sep_cols)  # 이미 데이터 있는 9월 date col 만
        prev_day_col = sep_data_cols[-1] if sep_data_cols else None
        sep_days_count = len(sep_data_cols)

        def vlookup(sku_ref: str, col_idx: int) -> str:
            return f"VLOOKUP({sku_ref},{MARGIN_LOOKUP_RANGE},{col_idx},FALSE)"

        def iferror(inner: str, fallback: str = '""') -> str:
            return f"IFERROR({inner},{fallback})"

        for r in range(data_start, max_row + 1):
            for c, kind in col_targets.items():
                # 원래 값이 없던 행(빈 행/서브총계 등)은 건너뜀
                orig = ws_v.cell(row=r, column=c).value
                if orig is None:
                    continue

                f = None
                if kind == '7avg':
                    rng = range_or_list(last7_cols, r)
                    if rng: f = f'=AVERAGE({rng})'
                elif kind == '14avg':
                    rng = range_or_list(last14_cols, r)
                    if rng: f = f'=AVERAGE({rng})'
                elif kind == '8avg':
                    rng = range_or_list(aug_cols, r)
                    if rng: f = f'=AVERAGE({rng})'
                elif kind == '8sum':
                    rng = range_or_list(aug_cols, r)
                    if rng: f = f'=SUM({rng})'
                elif kind == 'days' and stock_col and seven_avg_col:
                    j = f'{col_letter(stock_col)}{r}'
                    o = f'{col_letter(seven_avg_col)}{r}'
                    f = f'=IF({o}=0,0,{j}/{o})'
                elif kind == 'sepSum':
                    rng = range_or_list(sep_cols, r)
                    if rng: f = f'=SUM({rng})'
                elif kind == 'cost' and has_margin and sku_id_col:
                    sku_ref = f'{col_letter(sku_id_col)}{r}'
                    f = '=' + iferror(vlookup(sku_ref, MARGIN_COL_COST), '0')
                elif kind == 'costTotal' and cost_col and sep_sum_col:
                    f = f'={col_letter(cost_col)}{r}*{col_letter(sep_sum_col)}{r}'
                elif kind == 'supplyVat' and has_margin and sku_id_col:
                    sku_ref = f'{col_letter(sku_id_col)}{r}'
                    f = '=' + iferror(vlookup(sku_ref, MARGIN_COL_SUPPLY), '0')
                elif kind == 'supplyNoVat' and has_margin and sku_id_col:
                    sku_ref = f'{col_letter(sku_id_col)}{r}'
                    supply = vlookup(sku_ref, MARGIN_COL_SUPPLY)
                    vat = vlookup(sku_ref, MARGIN_COL_VAT)
                    f = '=' + iferror(f'{supply}-{vat}', '0')
                elif kind == 'margin' and has_margin and sku_id_col:
                    sku_ref = f'{col_letter(sku_id_col)}{r}'
                    f = '=' + iferror(vlookup(sku_ref, MARGIN_COL_MARGIN), '0')
                elif kind == 'marginTotal' and margin_col and sep_sum_col:
                    f = f'={col_letter(margin_col)}{r}*{col_letter(sep_sum_col)}{r}'
                elif kind == 'marginRate' and has_margin and sku_id_col:
                    sku_ref = f'{col_letter(sku_id_col)}{r}'
                    f = '=' + iferror(vlookup(sku_ref, MARGIN_COL_RATE), '0')
                elif kind == 'dailyRev' and prev_day_col and supply_novat_col:
                    f = f'={col_letter(prev_day_col)}{r}*{col_letter(supply_novat_col)}{r}'
                elif kind == 'totalRev' and sep_sum_col and supply_novat_col:
                    f = f'={col_letter(sep_sum_col)}{r}*{col_letter(supply_novat_col)}{r}'
                elif kind == 'avgDailyRev' and totalrev_col and sep_days_count > 0:
                    f = f'={col_letter(totalrev_col)}{r}/{sep_days_count}'
                if f:
                    ws.cell(row=r, column=c).value = f
                    formulas_added += 1

        # 숨김 + 서식 (hidden=True + width=0 → LibreOffice/Luckysheet 재변환에도 유지)
        for c in range(1, max_col + 1):
            letter = col_letter(c)
            if c in hidden:
                cd = ws.column_dimensions[letter]
                cd.hidden = True
                cd.width = 0
            fmt = decimals_format_for(headers.get(c, ''))
            if fmt:
                for r in range(1, max_row + 1):
                    cell = ws.cell(row=r, column=c)
                    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        cell.number_format = fmt

        print(f'  전처리 [{name}] hidden dates={len(hidden)} formulas→값={formula_replaced} 재삽입={formulas_added}')
    wb.save(dst)

    # 수식 결과값 캐시 (Luckysheet 가 즉시 화면에 표시)
    excel_recalc(dst)
    print('  재계산 + 캐시 저장')

    # 재계산이 hidden 상태를 지워버리므로 직접 XML 편집으로 재적용
    # (openpyxl 재저장은 LuckyExcel 파싱 오류 유발)
    apply_hidden_via_xml(dst)
    return dst


def apply_hidden_via_xml(xlsx_path: Path):
    """xlsx 를 zip 으로 열어 sheetN.xml 의 <cols> 를 직접 수정."""
    import zipfile, shutil as _sh, tempfile as _tf
    import xml.etree.ElementTree as ET
    import openpyxl

    # 시트별 hidden 컬럼 계산 (값 기반)
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    per_sheet_hidden: dict[str, set[int]] = {}
    for name in wb.sheetnames:
        ws = wb[name]
        max_col = ws.max_column or 1
        date_cols = detect_date_columns(ws, max_col)
        hidden = compute_hidden_date_cols(ws, date_cols)
        hidden |= ALWAYS_HIDDEN.get(name, set())
        per_sheet_hidden[name] = hidden

    # 시트 순서 (workbook.xml 의 sheetId 순으로 sheet1.xml, sheet2.xml, ...)
    # openpyxl sheetnames 는 표시 순서. xlsx 내부는 rId 순.
    # 안전하게 workbook.xml + xl/_rels 을 읽어 파일명 매핑
    with zipfile.ZipFile(xlsx_path, 'r') as z:
        wb_xml = z.read('xl/workbook.xml').decode('utf-8')
        rels_xml = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')

    # sheet name → rId
    name_to_rid = {}
    for m in re.finditer(r'<sheet\b[^/]*name="([^"]+)"[^/]*r:id="([^"]+)"', wb_xml):
        name_to_rid[m.group(1)] = m.group(2)
    # rId → target file (sheetN.xml)
    rid_to_file = {}
    for m in re.finditer(r'<Relationship\b[^/]*Id="([^"]+)"[^/]*Target="([^"]+)"', rels_xml):
        rid_to_file[m.group(1)] = m.group(2)

    tmp_dir = Path(_tf.mkdtemp())
    extracted = tmp_dir / 'contents'
    extracted.mkdir()
    with zipfile.ZipFile(xlsx_path, 'r') as z:
        z.extractall(extracted)

    for name, hidden in per_sheet_hidden.items():
        if not hidden: continue
        rid = name_to_rid.get(name)
        if not rid: continue
        rel = rid_to_file.get(rid, '')
        if not rel: continue
        sheet_path = extracted / 'xl' / rel.lstrip('/')
        if not sheet_path.exists(): continue
        text = sheet_path.read_text(encoding='utf-8')

        # 기존 <cols>...</cols> 를 파싱해 기존 폭/포맷 유지하면서 hidden 만 추가
        existing = []
        m_cols = re.search(r'<cols>(.*?)</cols>', text, flags=re.S)
        if m_cols:
            for m_col in re.finditer(r'<col\s+([^/]*?)/>', m_cols.group(1)):
                attrs = dict(re.findall(r'(\w+)="([^"]*)"', m_col.group(1)))
                try:
                    mn = int(attrs.get('min', '0')); mx = int(attrs.get('max', mn))
                except ValueError:
                    continue
                existing.append({'min': mn, 'max': mx, 'attrs': attrs})

        # 새 col 리스트 생성 (min..max range 를 hidden 여부에 따라 세분화)
        new_col_defs = []  # list of {min,max,attrs}
        # 우선 기존 range 를 hidden 컬럼과 non-hidden 으로 쪼갬
        covered = set()
        for e in existing:
            for c in range(e['min'], e['max']+1):
                covered.add(c)
                a = dict(e['attrs'])
                a['min'] = str(c); a['max'] = str(c)
                if c in hidden:
                    a['hidden'] = '1'
                    a['customWidth'] = '1'
                    a['width'] = '0'
                new_col_defs.append({'min': c, 'attrs': a})
        # 기존 <col> 에 없던 hidden 컬럼은 새로 추가
        for c in sorted(hidden):
            if c in covered: continue
            new_col_defs.append({'min': c, 'attrs': {'min': str(c), 'max': str(c), 'width': '0', 'hidden': '1', 'customWidth': '1'}})
        new_col_defs.sort(key=lambda x: x['min'])

        cols_body = ''.join(
            '<col ' + ' '.join(f'{k}="{v}"' for k, v in d['attrs'].items()) + '/>'
            for d in new_col_defs
        )
        new_cols = f'<cols>{cols_body}</cols>'
        if m_cols:
            text = text[:m_cols.start()] + new_cols + text[m_cols.end():]
        else:
            text = text.replace('<sheetData>', new_cols + '<sheetData>', 1)
        sheet_path.write_text(text, encoding='utf-8')
        print(f'  [{name}] XML hidden 적용 {len(hidden)}개')

    # 다시 zip 으로 묶음
    new_xlsx = tmp_dir / 'new.xlsx'
    with zipfile.ZipFile(new_xlsx, 'w', zipfile.ZIP_DEFLATED) as z:
        for root, _, files in os.walk(extracted):
            for f in files:
                fp = Path(root) / f
                arc = fp.relative_to(extracted).as_posix()
                z.write(fp, arc)
    _sh.move(str(new_xlsx), str(xlsx_path))
    _sh.rmtree(tmp_dir, ignore_errors=True)


HTML_TEMPLATE = r'''<!doctype html>
<html lang="ko">
<head>
<meta charset="utf-8">
<title>일일판매재고</title>
<link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/luckysheet@2.1.13/dist/plugins/css/pluginsCss.css">
<link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/luckysheet@2.1.13/dist/plugins/plugins.css">
<link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/luckysheet@2.1.13/dist/css/luckysheet.css">
<link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/luckysheet@2.1.13/dist/assets/iconfont/iconfont.css">
<style>
  html, body { margin: 0; padding: 0; height: 100%; overflow: hidden; font-family: "맑은 고딕", sans-serif; }
  #luckysheet { position: absolute; top: 0; left: 0; right: 0; bottom: 0; }
  #loader { position: fixed; top: 50%; left: 50%; transform: translate(-50%, -50%); font-size: 14px; color: #444; text-align: center; z-index: 10; }
  #loader .sp { display: inline-block; width: 20px; height: 20px; border: 3px solid #ddd; border-top-color: #2563eb; border-radius: 50%; animation: spin 1s linear infinite; vertical-align: middle; margin-right: 8px; }
  @keyframes spin { to { transform: rotate(360deg); } }
</style>
</head>
<body>
<div id="loader"><span class="sp"></span>엑셀 파싱 중...</div>
<div id="luckysheet"></div>
<script src="https://cdn.jsdelivr.net/npm/luckysheet@2.1.13/dist/plugins/js/plugin.js"></script>
<script src="https://cdn.jsdelivr.net/npm/luckysheet@2.1.13/dist/luckysheet.umd.js"></script>
<script src="https://cdn.jsdelivr.net/npm/luckyexcel@1.0.1/dist/luckyexcel.umd.js"></script>
<script>
const XLSX_B64 = "__XLSX_B64__";
// {sheetName: [0-indexed hidden col idx, ...]}
const HIDDEN_COLS = __HIDDEN_COLS_JSON__;

function b64ToUint8(b64) {
  const bin = atob(b64);
  const arr = new Uint8Array(bin.length);
  for (let i = 0; i < bin.length; i++) arr[i] = bin.charCodeAt(i);
  return arr;
}

function initSheet() {
  const loader = document.getElementById('loader');
  if (typeof LuckyExcel === 'undefined') {
    loader.textContent = 'LuckyExcel 로드 실패 (CDN 접근 불가?)';
    console.error('LuckyExcel undefined');
    return;
  }
  if (typeof luckysheet === 'undefined') {
    loader.textContent = 'Luckysheet 로드 실패';
    console.error('luckysheet undefined');
    return;
  }
  const bytes = b64ToUint8(XLSX_B64);
  const file = new File([bytes], 'daily.xlsx', {
    type: 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
  });
  console.log('file size:', file.size);
  LuckyExcel.transformExcelToLucky(file, (exportJson, luckysheetfile) => {
    console.log('transform callback', exportJson);
    if (!exportJson || !exportJson.sheets || !exportJson.sheets.length) {
      loader.textContent = '엑셀 파싱 실패 (콘솔 확인)';
      return;
    }
    // Hidden columns 를 Luckysheet 시트 config 에 주입
    for (const sh of exportJson.sheets) {
      const cols = HIDDEN_COLS[sh.name] || [];
      if (!cols.length) continue;
      sh.config = sh.config || {};
      sh.config.colhidden = sh.config.colhidden || {};
      for (const c of cols) sh.config.colhidden[c] = 0;
    }
    loader.style.display = 'none';
    luckysheet.create({
      container: 'luckysheet',
      data: exportJson.sheets,
      title: '일일판매재고',
      lang: 'en',
      showinfobar: false,
      showtoolbarConfig: {
        undoRedo: true, paintFormat: true, currencyFormat: true, percentageFormat: true,
        numberDecrease: true, numberIncrease: true, moreFormats: true,
        font: true, fontSize: true, bold: true, italic: true, strikethrough: true, underline: true,
        textColor: true, fillColor: true, border: true, mergeCell: true,
        horizontalAlignMode: true, verticalAlignMode: true, textWrapMode: true, textRotateMode: true,
        image: false, link: false, chart: false, postil: true, pivotTable: false,
        function: true, frozenMode: true, sortAndFilter: true, conditionalFormat: true, dataVerification: true,
        splitColumn: true, screenshot: false, findAndReplace: true, protection: true, print: false,
      },
      showsheetbarConfig: {
        add: true, menu: true, sheet: true,
      },
      cellRightClickConfig: {
        copy: true, copyAs: true, paste: true, insertRow: true, insertColumn: true,
        deleteRow: true, deleteColumn: true, deleteCell: true, hideRow: true, hideColumn: true,
        rowHeight: true, columnWidth: true, clear: true, matrix: true, sort: true, filter: true,
        chart: false, image: false, link: false, data: true, cellFormat: true,
      },
    });
  });
}
window.addEventListener('load', () => setTimeout(initSheet, 100));
window.addEventListener('error', (e) => {
  const l = document.getElementById('loader');
  if (l) l.innerHTML = '<b>스크립트 오류:</b><br>' + (e.message || e.error) + '<br><small>' + (e.filename || '') + ':' + (e.lineno || '') + '</small>';
});
</script>
</body>
</html>'''


def main():
    print(f'변환 중: {SRC_XLSB}')
    xlsx = xlsb_to_xlsx(SRC_XLSB)
    print(f'  → xlsx ({os.path.getsize(xlsx):,} bytes)')

    print('전처리 중...')
    prep = preprocess_xlsx(xlsx)
    print(f'  → {prep} ({os.path.getsize(prep):,} bytes)')

    data = prep.read_bytes()
    b64 = base64.b64encode(data).decode('ascii')
    print(f'base64 크기: {len(b64):,} chars')

    # 각 시트별 hidden 컬럼 (0-indexed) → JS 로 전달
    import json, openpyxl
    wb = openpyxl.load_workbook(prep, data_only=True)
    hidden_by_sheet = {}
    for name in wb.sheetnames:
        ws = wb[name]
        max_col = ws.max_column or 1
        date_cols = detect_date_columns(ws, max_col)
        hidden = compute_hidden_date_cols(ws, date_cols)
        hidden |= ALWAYS_HIDDEN.get(name, set())
        # 1-based → 0-based
        hidden_by_sheet[name] = sorted(c - 1 for c in hidden)
        print(f'  hidden [{name}]: {len(hidden)}개')

    hidden_json = json.dumps(hidden_by_sheet, ensure_ascii=False)

    html = HTML_TEMPLATE.replace('__XLSX_B64__', b64).replace('__HIDDEN_COLS_JSON__', hidden_json)
    OUT_HTML.write_text(html, encoding='utf-8')
    print(f'완료: {OUT_HTML}  ({os.path.getsize(OUT_HTML):,} bytes)')


if __name__ == '__main__':
    main()
